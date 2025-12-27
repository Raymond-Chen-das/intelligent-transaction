import os
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.utils import get_column_letter

def process_portfolio_analysis(excel_file, ic_sheet_name, factor_files):
    """
    處理投資組合分析
    
    Args:
        excel_file: 要處理的Excel檔案 (Real.xlsx, OLS.xlsx等)
        ic_sheet_name: IC工作表名稱 (真實IC, 預測IC等)
        factor_files: 因子檔案字典 {'bm': 'bm.xlsx', 'size': 'size.xlsx', 'mom': 'mom.xlsx'}
    """
    print(f"\n{'='*60}")
    print(f"處理檔案: {excel_file}")
    print(f"使用工作表: {ic_sheet_name}")
    print(f"{'='*60}")
    
    # 載入 Excel 檔案
    workbook = openpyxl.load_workbook(excel_file)
    sheet = workbook[ic_sheet_name]
    
    print("  📊 讀取IC數據...")
    # 讀取 IC 數據
    IC = pd.read_excel(excel_file, sheet_name=ic_sheet_name).T
    IC = np.array(IC.iloc[2:, 1:4])  # 取出 bm, size, mom 的 IC 值
    
    print("  📈 讀取因子數據...")
    # 讀取 bm 因子及報酬 (從 1999/01 開始,取 2013/12 之後)
    bm = pd.read_excel(factor_files['bm'], sheet_name="bm補值").T
    bm = np.array(bm.iloc[181:, :])  # 從第181行開始 (2013/12)
    bm_ret = pd.read_excel(factor_files['bm'], sheet_name="下個月月報酬補值").T
    bm_ret = np.array(bm_ret.iloc[181:, :])
    
    # 讀取 size 因子及報酬 (從 1999/01 開始,取 2013/12 之後)
    size = pd.read_excel(factor_files['size'], sheet_name="size補值").T
    size = np.array(size.iloc[181:, :])  # 從第181行開始 (2013/12)
    size_ret = pd.read_excel(factor_files['size'], sheet_name="下個月月報酬補值").T
    size_ret = np.array(size_ret.iloc[181:, :])
    
    # 讀取 mom 因子及報酬 (從 2000/01 開始,取 2013/12 之後)
    mom = pd.read_excel(factor_files['mom'], sheet_name="mom補值").T
    mom = np.array(mom.iloc[169:, :])  # 從第169行開始 (2013/12)
    mom_ret = pd.read_excel(factor_files['mom'], sheet_name="下個月月報酬補值").T
    mom_ret = np.array(mom_ret.iloc[169:, :])
    
    # 讀取公司ID (使用 bm 檔案)
    ID = pd.read_excel(factor_files['bm'], sheet_name="bm補值")
    ID = np.array(ID.iloc[:, 0:2])  # 取第1,2欄 (代號, 名稱)
    
    # ⚠️ 修正: 取所有資料中最小的時間點數量,避免索引超出範圍
    time_periods = min(IC.shape[0], bm.shape[0], size.shape[0], mom.shape[0])
    
    print(f"  ✓ 資料載入完成")
    print(f"    - IC 原始時間點: {IC.shape[0]}")
    print(f"    - bm 時間點: {bm.shape[0]}")
    print(f"    - size 時間點: {size.shape[0]}")
    print(f"    - mom 時間點: {mom.shape[0]}")
    print(f"    - 實際處理時間點: {time_periods} (取最小值)")
    print(f"    - 公司數量: {bm.shape[1]}")
    
    if IC.shape[0] > time_periods:
        print(f"  ⚠️  警告: IC資料有 {IC.shape[0]} 個時間點,但因子資料只有 {time_periods} 個")
        print(f"           將只處理前 {time_periods} 個時間點")
    
    print("\n  🔄 開始計算投資組合...")
    
    # 定義投資組合分組 (從 Row 12 開始)
    row_list = [12, 13, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23, 24]
    group_divisor = [10, 20, 50, 110, 120, 130, 150, 180, 200, 256, 400, 512, 1051]
    
    total_companies = bm.shape[1]  # 實際公司總數
    group_sizes = [round(total_companies / div) for div in group_divisor]
    
    for n in range(time_periods):
        # 進度顯示
        if (n + 1) % 20 == 0 or n == 0 or n == time_periods - 1:
            print(f"    進度: {n+1}/{time_periods} ({(n+1)/time_periods*100:.1f}%)")
        
        # 創建 DataFrame
        c = pd.DataFrame()
        sel_index = np.argmax(abs(IC[n]))  # 選擇絕對值最大的因子索引
        sel_value = IC[n][sel_index]
        
        # 根據選擇的因子設定對應的數據
        if sel_index == 0:
            fac, sel, ret = bm[n], 'bm', bm_ret[n]
        elif sel_index == 1:
            fac, sel, ret = size[n], 'size', size_ret[n]
        elif sel_index == 2:
            fac, sel, ret = mom[n], 'mom', mom_ret[n]
        
        c['factor'] = fac
        c['return'] = ret
        c['rank'] = c['factor'].rank(method='max')
        
        c = c.sort_values(by='rank')
        
        # 計算各組投資組合報酬
        for r, g in zip(row_list, group_sizes):
            high = np.mean(c['return'].iloc[-g:])  # 排名最高g個的平均報酬
            low = np.mean(c['return'].iloc[:g])    # 排名最低g個的平均報酬
            ans = high - low if sel_value > 0 else low - high
            sheet.cell(row=r, column=n+3, value=ans)
        
        # 將索引轉換成 list
        d = c.index.tolist()
        d1 = d[0]      # 排名最低的股票索引
        d2 = d[-1]     # 排名最高的股票索引
        a1 = ID[d1]    # 對應公司 ID/名稱
        a2 = ID[d2]    # 對應公司 ID/名稱
        
        # 寫入公司資料到Excel (Row 27-30)
        sheet.cell(row=27, column=n+3, value=a1[0])  # 低-股票代碼
        sheet.cell(row=28, column=n+3, value=a1[1])  # 低-公司名稱
        sheet.cell(row=29, column=n+3, value=a2[0])  # 高-股票代碼
        sheet.cell(row=30, column=n+3, value=a2[1])  # 高-公司名稱
    
    # 儲存檔案
    print(f"\n  💾 儲存檔案...")
    workbook.save(excel_file)
    print(f"  ✅ 完成!\n")


# ===== 主程式 =====
if __name__ == "__main__":
    # 設定因子檔案路徑
    factor_files = {
        'bm': 'bm.xlsx',
        'size': 'size.xlsx',
        'mom': 'mom.xlsx'
    }
    
    # 定義要處理的檔案及其對應的IC工作表名稱
    files_to_process = [
        ('Real.xlsx', '真實IC'),
        ('OLS.xlsx', '預測IC'),
        ('RF.xlsx', '預測IC'),
        ('NN3.xlsx', '預測IC'),
        ('NN4.xlsx', '預測IC'),
        ('NN5.xlsx', '預測IC'),
        ('XGBoost.xlsx', '預測IC'),
    ]
    
    print("\n" + "=" * 60)
    print("     批次處理投資組合分析系統")
    print("=" * 60)
    
    # 檢查因子檔案是否存在
    print("\n📋 檢查因子檔案...")
    all_files_exist = True
    for factor_name, factor_file in factor_files.items():
        if not os.path.exists(factor_file):
            print(f"  ❌ 錯誤: 找不到因子檔案 {factor_file}")
            all_files_exist = False
        else:
            print(f"  ✓ {factor_file}")
    
    if not all_files_exist:
        print("\n⚠️  請確保所有因子檔案都在程式同一目錄下!")
        exit(1)
    
    print("\n" + "=" * 60)
    
    # 處理每個檔案
    success_count = 0
    fail_count = 0
    
    for excel_file, ic_sheet_name in files_to_process:
        if not os.path.exists(excel_file):
            print(f"\n⚠️  警告: 找不到檔案 {excel_file}, 跳過...")
            fail_count += 1
            continue
        
        try:
            process_portfolio_analysis(excel_file, ic_sheet_name, factor_files)
            success_count += 1
        except Exception as e:
            print(f"\n❌ 錯誤: 處理 {excel_file} 時發生錯誤")
            print(f"   錯誤訊息: {str(e)}")
            import traceback
            traceback.print_exc()
            fail_count += 1
            continue
    
    # 顯示處理結果
    print("\n" + "=" * 60)
    print("     處理結果統計")
    print("=" * 60)
    print(f"  ✅ 成功: {success_count} 個檔案")
    print(f"  ❌ 失敗: {fail_count} 個檔案")
    print(f"  📊 總計: {success_count + fail_count} 個檔案")
    print("=" * 60)
    
    if success_count == len(files_to_process):
        print("\n🎉 所有檔案處理完成!")
    else:
        print("\n⚠️  部分檔案處理失敗,請檢查錯誤訊息")
    
    print()
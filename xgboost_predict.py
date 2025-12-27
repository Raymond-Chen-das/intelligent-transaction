""" Step1 引入需要的套件 """
import os
import openpyxl
import pandas as pd
import numpy as np
from openpyxl import load_workbook
from xgboost import XGBRegressor

""" Step2 定義路徑 """
os.chdir(r'C:\Users\raymo\Desktop\SCU work\智能交易')

""" Step3 建立模型 """
# XGBoost Regressor (保守、可重現)
def build_xgb():
    return XGBRegressor(
        n_estimators=100,
        max_depth=3,
        learning_rate=0.05,
        colsample_bytree=0.4,
        subsample=1.0,
        objective='reg:squarederror',
        random_state=999,
        n_jobs=1,
        verbosity=0
    )

""" Step4 預測IC值 """
# 讀取 Excel 檔案
workbook = openpyxl.load_workbook("XGBoost.xlsx")
sheet = workbook["預測IC"]  # 建模條件不存在，前先建立

# 預測 bm 的IC
df_x = pd.read_excel('bm.xlsx', sheet_name="bm補值")
df_y = pd.read_excel('bm.xlsx', sheet_name="bmIC")
df_x = df_x.T
df_y = df_y.T
df_x = df_x.iloc[2:,:]
df_y = df_y.iloc[2:,:]
n1, n2 = 178, 319
for n in range(n1, n2):
    X = df_x.iloc[0:n,:].values
    Y = df_y.iloc[0:n,:].values.ravel()
    test_X = df_x.iloc[n:n+1,:].values

    M = build_xgb()          # ☆ 每次重建模型 ☆
    M.fit(X, Y)
    predIC = M.predict(test_X)

    sheet.cell(row=3, column=n-n1+3, value=float(predIC[0]))
    print(n)

# 預測 size 的IC
df_x = pd.read_excel('size.xlsx', sheet_name="size補值")
df_y = pd.read_excel('size.xlsx', sheet_name="sizeIC")
df_x = df_x.T
df_y = df_y.T
df_x = df_x.iloc[2:,:]
df_y = df_y.iloc[2:,:]
n1, n2 = 178, 319
for n in range(n1, n2):
    X = df_x.iloc[0:n,:].values
    Y = df_y.iloc[0:n,:].values.ravel()
    test_X = df_x.iloc[n:n+1,:].values

    M = build_xgb()          # ☆ 每次重建模型 ☆
    M.fit(X, Y)
    predIC = M.predict(test_X)

    sheet.cell(row=4, column=n-n1+3, value=float(predIC[0]))
    print(n)

# 預測 mom 的IC
df_x = pd.read_excel('mom.xlsx', sheet_name="mom補值")
df_y = pd.read_excel('mom.xlsx', sheet_name="momIC")
df_x = df_x.T
df_y = df_y.T
df_x = df_x.iloc[2:,:]
df_y = df_y.iloc[2:,:]
n1, n2 = 166, 307
for n in range(n1, n2):
    X = df_x.iloc[0:n,:].values
    Y = df_y.iloc[0:n,:].values.ravel()
    test_X = df_x.iloc[n:n+1,:].values

    M = build_xgb()          # ☆ 每次重建模型 ☆
    M.fit(X, Y)
    predIC = M.predict(test_X)

    sheet.cell(row=5, column=n-n1+3, value=float(predIC[0]))
    print(n)

""" 存檔 """
workbook.save("XGBoost.xlsx")

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from copy import copy
import os

def copy_cell_format(source_cell, target_cell):
    """
    完整複製儲存格格式(包含字體、填充、邊框、對齊等)
    """
    if source_cell.has_style:
        target_cell.font = copy(source_cell.font)
        target_cell.fill = copy(source_cell.fill)
        target_cell.border = copy(source_cell.border)
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.number_format = copy(source_cell.number_format)
        target_cell.protection = copy(source_cell.protection)

def create_confusion_matrix_sheet(file_path, model_name):
    """
    為指定的 Excel 檔案創建混淆矩陣工作表
    
    Args:
        file_path: Excel 檔案路徑
        model_name: 模型名稱 (用於第一行標題)
    """
    print(f"處理檔案: {file_path}")
    
    # 載入 Excel 檔案
    wb = openpyxl.load_workbook(file_path)
    
    # 檢查是否已存在混淆矩陣工作表,若存在則刪除
    if '混淆矩陣' in wb.sheetnames:
        del wb['混淆矩陣']
        print(f"  - 刪除既有的「混淆矩陣」工作表")
    
    # 創建新的混淆矩陣工作表
    ws_new = wb.create_sheet('混淆矩陣')
    ws_rsquare = wb['R-square']
    
    # 獲取資料範圍 (C 到 EL, 即第3欄到第142欄)
    max_col = ws_rsquare.max_column
    
    print(f"  - 創建「混淆矩陣」工作表")
    print(f"  - 資料範圍: C 到 {get_column_letter(max_col)} (共 {max_col-2} 個時間點)")
    
    # ===== 第一部分: 複製預測值 (Row 1-5) =====
    # Row 1: 模型名稱
    ws_new['A1'] = model_name
    source_a1 = ws_rsquare['A1']
    copy_cell_format(source_a1, ws_new['A1'])
    
    # Row 2-5: 複製預測值資料及格式
    for row_idx in range(2, 6):  # Row 2-5
        for col_idx in range(1, max_col + 1):
            source_cell = ws_rsquare.cell(row=row_idx, column=col_idx)
            target_cell = ws_new.cell(row=row_idx, column=col_idx)
            target_cell.value = source_cell.value
            copy_cell_format(source_cell, target_cell)
    
    # Row 2: 添加「預測值」標籤
    ws_new['A2'] = '預測值'
    
    print(f"  - 複製預測值 (Row 1-5) 及格式")
    
    # ===== 第二部分: 複製真實值 (Row 7-11) =====
    # Row 7: Real
    ws_new['A7'] = 'Real'
    source_a7 = ws_rsquare['A7']
    copy_cell_format(source_a7, ws_new['A7'])
    
    # Row 8-11: 複製真實值資料及格式
    for row_idx in range(8, 12):  # Row 8-11
        for col_idx in range(1, max_col + 1):
            source_cell = ws_rsquare.cell(row=row_idx, column=col_idx)
            target_cell = ws_new.cell(row=row_idx, column=col_idx)
            target_cell.value = source_cell.value
            copy_cell_format(source_cell, target_cell)
    
    # Row 8: 添加「真實值」標籤
    ws_new['A8'] = '真實值'
    
    print(f"  - 複製真實值 (Row 7-11) 及格式")
    
    # 取得參考格式 (從 R-square 工作表)
    ref_label_cell = ws_rsquare['A13']  # 參考標籤格式
    ref_data_cell = ws_rsquare['C13']   # 參考數據格式
    
    # ===== 第三部分: 創建 TP 區塊 (Row 13-16) =====
    ws_new['A13'] = 'TP'
    ws_new['A14'] = '預測>0,真實>0'
    copy_cell_format(ref_label_cell, ws_new['A13'])
    copy_cell_format(ref_label_cell, ws_new['A14'])
    
    ws_new['B13'] = 'bm'
    ws_new['B14'] = 'size'
    ws_new['B15'] = 'mom'
    ws_new['B16'] = 'TP'
    
    # 複製 B 欄的格式
    for row in [13, 14, 15, 16]:
        source_b = ws_rsquare.cell(row=row-10, column=2)  # 參考 Row 3-6 的 B 欄
        copy_cell_format(source_b, ws_new.cell(row=row, column=2))
    
    # 填入 TP 公式及格式
    for col_idx in range(3, max_col + 1):  # 從 C 欄開始
        col_letter = get_column_letter(col_idx)
        
        # bm 的 TP
        ws_new[f'{col_letter}13'] = f'=IF(AND({col_letter}3>0,{col_letter}9>0),1,0)'
        copy_cell_format(ref_data_cell, ws_new[f'{col_letter}13'])
        
        # size 的 TP
        ws_new[f'{col_letter}14'] = f'=IF(AND({col_letter}4>0,{col_letter}10>0),1,0)'
        copy_cell_format(ref_data_cell, ws_new[f'{col_letter}14'])
        
        # mom 的 TP
        ws_new[f'{col_letter}15'] = f'=IF(AND({col_letter}5>0,{col_letter}11>0),1,0)'
        copy_cell_format(ref_data_cell, ws_new[f'{col_letter}15'])
    
    # TP 總計公式 - 使用更簡潔的寫法
    sum_parts = []
    for col_idx in range(3, max_col + 1):
        col_letter = get_column_letter(col_idx)
        sum_parts.append(f'SUM({col_letter}13:{col_letter}15)')
    ws_new['C16'] = f'={"+".join(sum_parts)}'
    copy_cell_format(ref_data_cell, ws_new['C16'])
    
    # TP 字體設為藍色
    ws_new['B16'].font = Font(name=ws_new['B16'].font.name, 
                               size=ws_new['B16'].font.size, 
                               color='0000FF', 
                               bold=ws_new['B16'].font.bold)
    
    print(f"  - 創建 TP 區塊 (Row 13-16)")
    
    # ===== 第四部分: 創建 FP 區塊 (Row 17-20) =====
    ws_new['A17'] = 'FP'
    ws_new['A18'] = '預測>0,真實<0'
    copy_cell_format(ref_label_cell, ws_new['A17'])
    copy_cell_format(ref_label_cell, ws_new['A18'])
    
    ws_new['B17'] = 'bm'
    ws_new['B18'] = 'size'
    ws_new['B19'] = 'mom'
    ws_new['B20'] = 'FP'
    
    for row in [17, 18, 19, 20]:
        source_b = ws_rsquare.cell(row=row-14, column=2)
        copy_cell_format(source_b, ws_new.cell(row=row, column=2))
    
    # 填入 FP 公式及格式
    for col_idx in range(3, max_col + 1):
        col_letter = get_column_letter(col_idx)
        ws_new[f'{col_letter}17'] = f'=IF(AND({col_letter}3>0,{col_letter}9<0),1,0)'
        ws_new[f'{col_letter}18'] = f'=IF(AND({col_letter}4>0,{col_letter}10<0),1,0)'
        ws_new[f'{col_letter}19'] = f'=IF(AND({col_letter}5>0,{col_letter}11<0),1,0)'
        
        for row in [17, 18, 19]:
            copy_cell_format(ref_data_cell, ws_new[f'{col_letter}{row}'])
    
    # FP 總計公式
    sum_parts = []
    for col_idx in range(3, max_col + 1):
        col_letter = get_column_letter(col_idx)
        sum_parts.append(f'SUM({col_letter}17:{col_letter}19)')
    ws_new['C20'] = f'={"+".join(sum_parts)}'
    copy_cell_format(ref_data_cell, ws_new['C20'])
    
    ws_new['B20'].font = Font(name=ws_new['B20'].font.name, 
                               size=ws_new['B20'].font.size, 
                               color='0000FF', 
                               bold=ws_new['B20'].font.bold)
    
    print(f"  - 創建 FP 區塊 (Row 17-20)")
    
    # ===== 第五部分: 創建 FN 區塊 (Row 21-24) =====
    ws_new['A21'] = 'FN'
    ws_new['A22'] = '預測<0,真實>0'
    copy_cell_format(ref_label_cell, ws_new['A21'])
    copy_cell_format(ref_label_cell, ws_new['A22'])
    
    ws_new['B21'] = 'bm'
    ws_new['B22'] = 'size'
    ws_new['B23'] = 'mom'
    ws_new['B24'] = 'FN'
    
    for row in [21, 22, 23, 24]:
        source_b = ws_rsquare.cell(row=row-18, column=2)
        copy_cell_format(source_b, ws_new.cell(row=row, column=2))
    
    # 填入 FN 公式及格式
    for col_idx in range(3, max_col + 1):
        col_letter = get_column_letter(col_idx)
        ws_new[f'{col_letter}21'] = f'=IF(AND({col_letter}3<0,{col_letter}9>0),1,0)'
        ws_new[f'{col_letter}22'] = f'=IF(AND({col_letter}4<0,{col_letter}10>0),1,0)'
        ws_new[f'{col_letter}23'] = f'=IF(AND({col_letter}5<0,{col_letter}11>0),1,0)'
        
        for row in [21, 22, 23]:
            copy_cell_format(ref_data_cell, ws_new[f'{col_letter}{row}'])
    
    # FN 總計公式
    sum_parts = []
    for col_idx in range(3, max_col + 1):
        col_letter = get_column_letter(col_idx)
        sum_parts.append(f'SUM({col_letter}21:{col_letter}23)')
    ws_new['C24'] = f'={"+".join(sum_parts)}'
    copy_cell_format(ref_data_cell, ws_new['C24'])
    
    ws_new['B24'].font = Font(name=ws_new['B24'].font.name, 
                               size=ws_new['B24'].font.size, 
                               color='0000FF', 
                               bold=ws_new['B24'].font.bold)
    
    print(f"  - 創建 FN 區塊 (Row 21-24)")
    
    # ===== 第六部分: 創建 TN 區塊 (Row 25-28) =====
    ws_new['A25'] = 'TN'
    ws_new['A26'] = '預測<0,真實<0'
    copy_cell_format(ref_label_cell, ws_new['A25'])
    copy_cell_format(ref_label_cell, ws_new['A26'])
    
    ws_new['B25'] = 'bm'
    ws_new['B26'] = 'size'
    ws_new['B27'] = 'mom'
    ws_new['B28'] = 'TN'
    
    for row in [25, 26, 27, 28]:
        source_b = ws_rsquare.cell(row=row-22, column=2)
        copy_cell_format(source_b, ws_new.cell(row=row, column=2))
    
    # 填入 TN 公式及格式
    for col_idx in range(3, max_col + 1):
        col_letter = get_column_letter(col_idx)
        ws_new[f'{col_letter}25'] = f'=IF(AND({col_letter}3<0,{col_letter}9<0),1,0)'
        ws_new[f'{col_letter}26'] = f'=IF(AND({col_letter}4<0,{col_letter}10<0),1,0)'
        ws_new[f'{col_letter}27'] = f'=IF(AND({col_letter}5<0,{col_letter}11<0),1,0)'
        
        for row in [25, 26, 27]:
            copy_cell_format(ref_data_cell, ws_new[f'{col_letter}{row}'])
    
    # TN 總計公式
    sum_parts = []
    for col_idx in range(3, max_col + 1):
        col_letter = get_column_letter(col_idx)
        sum_parts.append(f'SUM({col_letter}25:{col_letter}27)')
    ws_new['C28'] = f'={"+".join(sum_parts)}'
    copy_cell_format(ref_data_cell, ws_new['C28'])
    
    ws_new['B28'].font = Font(name=ws_new['B28'].font.name, 
                               size=ws_new['B28'].font.size, 
                               color='0000FF', 
                               bold=ws_new['B28'].font.bold)
    
    print(f"  - 創建 TN 區塊 (Row 25-28)")
    
    # ===== 第七部分: 創建評估指標 (Row 30-33) =====
    ws_new['A30'] = 'Accuracy'
    ws_new['B30'] = '準確率'
    ws_new['C30'] = '=(C16+C28)/(C16+C28+C20+C24)'
    
    ws_new['A31'] = 'Precision'
    ws_new['B31'] = '精確率'
    ws_new['C31'] = '=C16/(C16+C20)'
    
    ws_new['A32'] = 'Recall'
    ws_new['B32'] = '召回率'
    ws_new['C32'] = '=C16/(C16+C24)'
    
    ws_new['A33'] = 'F1-score'
    ws_new['B33'] = 'F1分數'
    ws_new['C33'] = '=2*(C31*C32)/(C31+C32)'
    
    # 複製評估指標的基本格式
    for row in [30, 31, 32, 33]:
        # A 欄和 B 欄使用標籤格式
        copy_cell_format(ref_label_cell, ws_new[f'A{row}'])
        copy_cell_format(ref_label_cell, ws_new[f'B{row}'])
        # C 欄使用數據格式
        copy_cell_format(ref_data_cell, ws_new[f'C{row}'])
    
    # 設定評估指標數值為紅色
    for row in [30, 31, 32, 33]:
        ws_new[f'C{row}'].font = Font(name=ws_new[f'C{row}'].font.name,
                                       size=ws_new[f'C{row}'].font.size,
                                       color='FF0000',
                                       bold=ws_new[f'C{row}'].font.bold)
    
    print(f"  - 創建評估指標 (Row 30-33)")
    
    # 儲存檔案
    wb.save(file_path)
    print(f"  ✓ 完成! 檔案已儲存\n")


# ===== 主程式 =====
if __name__ == "__main__":
    # 定義檔案清單及對應的模型名稱
    files = [
        ('OLS.xlsx', 'OLS'),
        ('RF.xlsx', 'RF'),
        ('NN3.xlsx', 'NN3'),
        ('NN4.xlsx', 'NN4'),
        ('NN5.xlsx', 'NN5')
    ]
    
    print("=" * 60)
    print("開始批次處理混淆矩陣工作表")
    print("=" * 60 + "\n")
    
    for file_name, model_name in files:
        # 檢查檔案是否存在
        if not os.path.exists(file_name):
            print(f"⚠️  警告: 找不到檔案 {file_name},跳過...")
            continue
        
        try:
            create_confusion_matrix_sheet(file_name, model_name)
        except Exception as e:
            print(f"  ✗ 錯誤: {str(e)}\n")
            continue
    
    print("=" * 60)
    print("所有檔案處理完成!")
    print("=" * 60)